import re

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
import openpyxl
import pypdf
import pdfplumber

from .pdf_table import (bat_dong_so_cot, checksum_gap, column_names,
                        hang_khong_gia_tri, merge_table_rows, row_to_text,
                        split_header_body)
from .xlsx_header import compose_two_tier, find_header
from src.ocr import table
from src.ocr.document import read_page
from src.ocr.engine import TesseractMissing

# Nhánh số CHỈ nhận numbering đa cấp ("1.1", "3.2.1"), sub-level 1-2 chữ số:
# numbering 1 cấp ("1. ...") là KHOẢN (nội dung) trong luật VN chứ không phải
# heading, còn nhóm 3 chữ số là dấu phân cách nghìn ("5.000.000.000 đồng.")
# hoặc mã HS phụ lục ("2931.9080") — cả ba từng bị nhận nhầm heading khiến
# chunking nuốt nội dung (spec 2026-07-15-rag-heading-detection-fix).
#
# Nhánh `Điều` SIẾT LẠI 2026-08-19 (spec ingest-hygiene §4): trước đây là
# `^\s*(Chương|Mục|Điều)\b`, nên khớp cả tham chiếu chéo GIỮA câu — "Điều 11
# của Luật này quy định.", "Điều ước quốc tế mà Cộng hòa..." — sinh ra 15 mục
# là mảnh câu, mỗi mảnh CẮT ĐÔI một Điều thật. Đây là lỗi ingest DUY NHẤT đã
# chứng minh gây hại đo được: chuỗi "Điều ước quốc tế..." chiếm hạng 1 của một
# câu hỏi thật (spec P0 §11.1(b)). Nay `Điều` phải mang dạng "Điều <số>." rồi
# mới tới tiêu đề.
#
# `Chương` và `Mục` GIỮ NGUYÊN: chúng sinh 190 mục trống, nhưng đó là việc của
# P3b (phân cấp) — đụng vào đây là trộn thêm một biến vào cùng một lần re-index.
_HEADING_RE = re.compile(
    r"^\s*(Chương|Mục)\b"
    r"|^\s*Điều\s+\d+\s*[\.\-–]\s*\S"
    r"|^\s*\d+\.\d{1,2}(\.\d{1,2})*(?!\d)[\.\)]?\s+\S")

_CHUONG_RE = re.compile(r"^\s*Chương\b")
_MUC_RE = re.compile(r"^\s*Mục\b")
_DIEU_RE = re.compile(r"^\s*Điều\s+\d+\s*[\.\-–]\s*\S")


def heading_level(text: str) -> int | None:
    """Cấp của một dòng tiêu đề, hoặc None nếu không phải tiêu đề.

    TRƯỚC 2026-08-20 mọi tiêu đề đều là cấp 2, tức `Chương`, `Mục` và `Điều`
    là ANH EM. Vì `chunk_text_blocks` đẩy khỏi stack mọi mục cùng-hoặc-cao-cấp,
    "Điều 102" hất "Chương VI" ra và section_path thành PHẮNG.

    Tác hại đo được: 27 cặp điều luật trong corpus TRÙNG TIÊU ĐỀ từng chữ, và
    thứ phân biệt chúng nằm ở chương. Ví dụ đã bắt được trong bộ synthesis_live:

        Điều 70  ← Chương V  — BẢO HIỂM XÃ HỘI BẮT BUỘC
        Điều 102 ← Chương VI — BẢO HIỂM XÃ HỘI TỰ NGUYỆN

    cả hai cùng tên "Hưởng bảo hiểm xã hội một lần". Hỏi về TỰ NGUYỆN thì cả
    dense lẫn cross-encoder đều chọn Điều 70, vì từ "tự nguyện" không dính vào
    văn bản index của Điều 102 — trớ trêu là nó lại nằm trong THÂN Điều 70
    ("không bao gồm số tiền ngân sách nhà nước hỗ trợ đóng bảo hiểm xã hội tự
    nguyện"). Đáp án đúng không lọt nổi top-6.

    THANG CẤP và lý do khoảng cách:

        1  Chương
        2  dòng IN HOA          — tiêu đề của chương nằm ở DÒNG SAU số chương
                                  ("Chương VI" / "BẢO HIỂM XÃ HỘI TỰ NGUYỆN"),
                                  và chính nó mang từ phân biệt. Phải nằm GIỮA
                                  Chương và Mục: để cùng cấp với Mục thì Mục đầu
                                  tiên sẽ hất nó ra và mất đúng từ đó.
        3  Mục
        4  Điều
        5  numbering đa cấp

    Dòng IN HOA ở cấp 2 còn là hàng rào an toàn: một dòng IN HOA lạc giữa
    chương (vd "ĐIỀU KHOẢN THI HÀNH" in giữa văn bản) chỉ hất được các Mục,
    KHÔNG hất được Chương — nên chương vẫn sống sót cho mọi điều phía sau.
    """
    if _CHUONG_RE.match(text):
        return 1
    if _MUC_RE.match(text):
        return 3
    if _DIEU_RE.match(text):
        return 4
    if _HEADING_RE.match(text):
        return 5      # numbering đa cấp "1.1", "3.2.1"
    if text.isupper() and len(text) <= 80:
        return 2
    return None


# ---------------------------------------------------------------------------
# Suy phân cấp cho .docx — spec 2026-09-04
#
# Vì sao KHÔNG nới `heading_level()` ở trên: nó đang phục vụ corpus LUẬT và đã
# hiệu chỉnh kỹ (nhánh `Điều` từng bị siết vì sinh 15 mục là mảnh câu). Nới nó
# ra là rủi ro hồi quy trên đường đang chạy tốt, để chữa một đường khác. Ở đây
# dùng LẠI các regex của nó nhưng xếp thứ tự kiểm khác và thêm mẫu hành chính.
#
# Thang RỘNG HƠN thang 1-5 của `heading_level()` để chèn được các tầng hành
# chính vào giữa. `chunk_text_blocks` chỉ so cấp bằng `>=` nên thang không cần
# liền mạch. Số ở đây là ĐỀ XUẤT, hiệu chỉnh ở Task 4.
DOCX_LEVEL = {
    "phan": 5,
    "chuong": 10,
    "upper": 20,
    "muc": 30,
    "roman": 35,
    "letter": 38,
    "dieu": 40,
    "arabic": 45,
    "multi": 50,
}

# Cấp của style Word (`Heading 1..9`) nhân với hệ số này để về CÙNG thang.
# Dùng thô là sai: `Heading 2` sẽ thành cấp 2, cao hơn cả `phan` (5) lẫn
# `chuong` (10), và hất sạch mọi thứ phía trên nó.
STYLE_SCALE = 10

_PHAN_RE = re.compile(r"^\s*PHẦN\s+\S", re.IGNORECASE)

# Một dòng mở đầu bằng ký hiệu đánh số TRẦN: "I.", "A.", "12.", "3)".
# `\s+\S` bắt buộc có khoảng trắng rồi mới tới nội dung — nhờ đó
# "5.000.000 đồng" không khớp (sau dấu chấm là chữ số, không phải khoảng trắng).
_ENUM_RE = re.compile(r"^\s*([A-ZĐ]+|\d{1,2})\s*[.)]\s+\S")

# Mục đánh số ĐA CẤP: "12.1", "3.2.1". Dùng để lấy TIỀN TỐ làm bằng chứng cho
# mục cha ("12." là tiêu đề vì "12.1" tồn tại).
_MULTI_RE = re.compile(r"^\s*(\d{1,2})\.\d{1,2}")

_ROMAN_VALUE = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100, "D": 500, "M": 1000}


def _roman_to_int(token: str) -> int | None:
    """Giá trị của một số La Mã, hoặc None nếu không phải số La Mã."""
    if not token or any(c not in _ROMAN_VALUE for c in token):
        return None
    total = 0
    highest = 0
    for char in reversed(token):
        value = _ROMAN_VALUE[char]
        total = total - value if value < highest else total + value
        highest = max(highest, value)
    return total or None


def _has_consecutive(numbers: set[int]) -> bool:
    """Tập có chứa hai số KẾ TIẾP NHAU không (n và n+1).

    Đây là "bằng chứng" của một họ đánh số: một tài liệu có `I.` rồi `II.` thì
    họ La Mã là thật; một `I.` đứng một mình có thể chỉ là chữ cái lạc.
    Nói về THỨ TỰ ĐÁNH SỐ, không phải khoảng cách dòng."""
    return any(n + 1 in numbers for n in numbers)


def _collect_evidence(texts: list[str]) -> tuple[bool, bool, bool, set[int]]:
    """Bằng chứng đánh số mà CHÍNH tài liệu đưa ra.

    Trả `(roman_ok, letter_ok, arabic_ok, multi_parents)`.

    Một token một ký tự như "I" được tính vào CẢ họ La Mã lẫn họ chữ cái; việc
    nó thuộc họ nào do bằng chứng quyết định sau, không đoán trước.
    """
    romans: set[int] = set()
    letters: set[int] = set()
    arabics: set[int] = set()
    parents: set[int] = set()

    for text in texts:
        multi = _MULTI_RE.match(text)
        if multi:
            parents.add(int(multi.group(1)))
            continue
        found = _ENUM_RE.match(text)
        if not found:
            continue
        token = found.group(1)
        if token.isdigit():
            arabics.add(int(token))
            continue
        value = _roman_to_int(token)
        if value is not None:
            romans.add(value)
        if len(token) == 1:
            letters.add(ord(token))

    return (_has_consecutive(romans), _has_consecutive(letters),
            _has_consecutive(arabics), parents)


def _docx_level(text: str, roman_ok: bool, letter_ok: bool,
                arabic_ok: bool, parents: set[int]) -> int | None:
    """Cấp của MỘT dòng, với bằng chứng của cả tài liệu đã tính sẵn.

    THỨ TỰ KIỂM quan trọng: nhánh đánh số phải đứng TRƯỚC nhánh dòng IN HOA
    của `heading_level()`. Nếu không, "II. CÁ NHÂN CƯ TRÚ" (in hoa) ra cấp
    khác "I. Đặc điểm hoạt động" (thường) dù cùng một họ, và stack breadcrumb
    lồng sai.
    """
    found = _ENUM_RE.match(text)
    if found:
        token = found.group(1)
        if token.isdigit():
            number = int(token)
            if number in parents or arabic_ok:
                return DOCX_LEVEL["arabic"]
            return None
        if roman_ok and _roman_to_int(token) is not None:
            return DOCX_LEVEL["roman"]
        if letter_ok and len(token) == 1:
            return DOCX_LEVEL["letter"]
        # Có đánh số nhưng tài liệu không đưa ra bằng chứng nào — rơi xuống
        # các nhánh dưới thay vì đoán bừa.

    if _PHAN_RE.match(text):
        return DOCX_LEVEL["phan"]

    shared = heading_level(text)
    if shared is None:
        return None
    return {1: DOCX_LEVEL["chuong"], 2: DOCX_LEVEL["upper"],
            3: DOCX_LEVEL["muc"], 4: DOCX_LEVEL["dieu"],
            5: DOCX_LEVEL["multi"]}[shared]


def docx_heading_levels(texts: list[str]) -> list[int | None]:
    """Cấp tiêu đề cho từng dòng của MỘT tài liệu .docx, hoặc None.

    Nhận cả tài liệu chứ không chấm từng dòng độc lập, vì quy tắc "đánh số
    trần phải tự chứng minh" (spec mục 5) cần bằng chứng ở phạm vi tài liệu.
    Nhận `list[str]` chứ KHÔNG nhận đối tượng docx — module giữ nguyên tính
    lá thuần, và test không cần tệp thật.
    """
    roman_ok, letter_ok, arabic_ok, parents = _collect_evidence(texts)
    return [_docx_level(t, roman_ok, letter_ok, arabic_ok, parents)
            for t in texts]


_DIGITS_RE = re.compile(r"\d+")


def _normalize_digits(text: str) -> str:
    """Thay mọi chuỗi chữ số bằng '#'.

    Không có bước này thì bộ lọc tần suất bỏ sót hoàn toàn: 'about:blank 5/164'
    và 'about:blank 6/164' là hai chuỗi khác nhau, mỗi cái chỉ xuất hiện đúng
    một trang."""
    return _DIGITS_RE.sub("#", text)


def detect_page_furniture(pages: list[list[str]], *, min_pages: int = 5,
                          page_ratio: float = 0.6, edge_ratio: float = 0.9,
                          edge: int = 2) -> set[str]:
    """Dạng-đã-chuẩn-hoá của các dòng là header/footer trang.

    Điều kiện KÉP, không phải hoặc (spec 2026-08-19-ingest-hygiene §3):
      (a) xuất hiện trên >= page_ratio số trang, VÀ
      (b) >= edge_ratio số lần xuất hiện nằm trong `edge` dòng đầu hoặc cuối
          của trang.

    Vì sao cần (b): chuẩn hoá chữ số gộp NHIỀU hàng bảng khác nhau thành một
    nhóm. Đo thật trên luat-thuexuatnhapkhau.pdf: nhóm '# #.#' (hàng bảng mã
    HS) đạt 48% số trang — chỉ thoát ngưỡng 60% nhờ 12 điểm phần trăm, quá
    mỏng để tin. Với (b) thì nó ra 15% và bị loại dứt khoát, trong khi 18 nhóm
    rác thật của 9 tệp đều đạt 100%/100%.

    Vì sao không khớp mẫu 'about:blank': đó là dấu vết của CÁCH IN bộ PDF này
    (print-to-PDF từ trình duyệt). Tài liệu nguồn khác mang rác khác; tần suất
    bắt được cả loại chưa gặp, khớp mẫu cứng thì không.
    """
    n_pages = len(pages)
    if n_pages < min_pages:
        return set()
    seen_pages: dict[str, set[int]] = {}
    at_edge: dict[str, int] = {}
    total: dict[str, int] = {}
    for pageno, lines in enumerate(pages):
        for i, text in enumerate(lines):
            key = _normalize_digits(text)
            seen_pages.setdefault(key, set()).add(pageno)
            total[key] = total.get(key, 0) + 1
            if i < edge or i >= len(lines) - edge:
                at_edge[key] = at_edge.get(key, 0) + 1
    return {key for key, pgs in seen_pages.items()
            if len(pgs) / n_pages >= page_ratio
            and at_edge.get(key, 0) / total[key] >= edge_ratio}


# "hiệu lực THI HÀNH từ ngày ..." — KHÔNG phải "có hiệu lực từ ngày ...".
# Cụm "có hiệu lực" xuất hiện dày đặc trong NỘI DUNG điều luật (hiệu lực của
# hợp đồng, của giao dịch dân sự), và bản đầu của regex này khớp đúng chúng:
# đo trên 9 PDF thì 8 tệp trả về câu về hiệu lực HỢP ĐỒNG, chỉ 1 tệp ra ngày
# thật. Chữ "thi hành" là thứ phân biệt ngày của CHÍNH VĂN BẢN với mọi cách
# dùng khác.
_EFFECTIVE_RE = re.compile(
    r"hiệu\s+lực\s+thi\s+hành\s+(?:kể\s+)?từ\s+ngày\s+"
    r"(\d{1,2})\s+tháng\s+(\d{1,2})\s+năm\s+(\d{4})",
    re.IGNORECASE)


def extract_effective_date(text: str | None):
    """Ngày hiệu lực của văn bản, hoặc None.

    None là kết quả HỢP LỆ, không phải lỗi: 8 tài liệu nghiệp vụ (.docx,
    .xlsx) trong corpus không phải văn bản quy phạm và không có ngày hiệu lực.
    Cả 9/9 PDF luật thì đều đọc được ngày (đo 2026-08-20 trên DB thật).

    Con số 9/9 này SỬA LẠI một phép đo sai của chính tôi. Phép đo đầu báo 8/9
    vì nó đi tìm MỤC mang tên "Hiệu lực thi hành"; `luat-doanhnghiep.pdf` đặt
    tên mục là "Điều khoản thi hành" nên bị bỏ sót, dù câu "Luật này có hiệu
    lực thi hành từ ngày 01 tháng 01 năm 2021" nằm ngay trong đó. Tìm theo CÂU
    trên toàn văn không phụ thuộc vào cách đặt tên mục, nên bắt được cả hai.

    Ngày không hợp lệ (45 tháng 13) trả None thay vì ném: một cú trích hỏng
    không được làm vỡ việc ingest cả tài liệu."""
    import datetime as _dt
    if not text:
        return None
    m = _EFFECTIVE_RE.search(text)
    if not m:
        return None
    day, month, year = (int(x) for x in m.groups())
    try:
        return _dt.date(year, month, day)
    except ValueError:
        return None


def _bang_thanh_text(tbl) -> str:
    """Một bảng .docx thành text nhiều dòng, mỗi dòng một hàng, cột ngăn bởi "|".

    Giữ HÀNG TIÊU ĐỀ ở dòng đầu vì đó là thứ mang nghĩa cho mọi hàng dưới: một
    ô "5%" tách khỏi cột "Chiết khấu" là vô nghĩa.

    Ô gộp (merged) làm `row.cells` lặp lại cùng nội dung — chấp nhận, vì lặp
    một nhãn còn hơn mất nó.
    """
    dong = []
    for hang in tbl.rows:
        o = [c.text.strip().replace("\n", " ") for c in hang.cells]
        if any(o):
            dong.append(" | ".join(o))
    return "\n".join(dong)


def parse_docx(path: str) -> list[dict]:
    """Blocks in order; a heading block carries heading_level (1..n), body carries None.

    DUYỆT THÂN TÀI LIỆU THEO THỨ TỰ, không phải `doc.paragraphs` rồi `doc.tables`.
    Trước 2026-08-22 hàm này chỉ đọc `doc.paragraphs`, tức **bỏ 100% bảng biểu** —
    bảng chính sách chiết khấu, SLA, định mức bị nuốt sạch. Tệ hơn: tệp vẫn sinh
    chunk từ các đoạn văn nên `IngestError` KHÔNG kích hoạt, và trợ lý tự tin trả
    lời "chính sách không quy định" cho thứ có trong tài liệu.

    Vì sao không thể chỉ nối thêm `for t in doc.tables`: `chunk_text_blocks` dựng
    `section_path` theo THỨ TỰ block, nên dồn mọi bảng xuống cuối sẽ gắn chúng vào
    chương SAI. Bảng bị gắn nhầm chương còn nguy hiểm hơn bảng bị bỏ — nó trả lời
    sai một cách tự tin thay vì im lặng.

    `doc.element.body` là nơi duy nhất giữ đúng thứ tự xen kẽ đoạn văn / bảng;
    `doc.paragraphs` cũng KHÔNG chứa đoạn nằm bên trong ô bảng.

    HAI LƯỢT từ 2026-09-04: gom tất cả text trước (chưa quyết định cấp), suy cấp
    từ mẫu chữ của toàn tài liệu, rồi mới dựng block. Một lượt thì không thể biết
    một mẫu "(I., II., III.)" hay chỉ lạc chữ cái, cần bằng chứng ở phạm vi tài
    liệu (spec 2026-09-04 mục 7). Style Heading được ánh xạ sang thang DOCX_LEVEL
    bằng `N * STYLE_SCALE`, không dùng cấp thô.
    """
    doc = Document(path)

    # Lượt 1: gom theo ĐÚNG thứ tự thân tài liệu, chưa quyết định cấp.
    items: list[tuple[str, str, int | None]] = []   # (kind, text, style_level)
    for child in doc.element.body.iterchildren():
        tag = etree.QName(child).localname if hasattr(child, "tag") else ""
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "") if para.style else ""
            style_level = None
            if style.startswith("Heading"):
                try:
                    style_level = int(style.split()[-1]) * STYLE_SCALE
                except ValueError:
                    style_level = DOCX_LEVEL["chuong"]
            items.append(("p", text, style_level))
        elif tag == "tbl":
            text = _bang_thanh_text(Table(child, doc))
            if text:
                items.append(("tbl", text, None))

    # Lượt 2: suy cấp từ chữ, CHỈ trên đoạn văn. Không đưa text bảng vào: một
    # ô bảng chứa "1." sẽ làm nhiễu bằng chứng đánh số của cả tài liệu.
    para_levels = docx_heading_levels([t for kind, t, _ in items if kind == "p"])

    blocks: list[dict] = []
    para_index = 0
    for kind, text, style_level in items:
        if kind == "tbl":
            # heading_level=None: bảng là THÂN, không bao giờ là tiêu đề —
            # để nó thành heading sẽ phá breadcrumb của cả mục.
            blocks.append({"text": text, "heading_level": None, "page": None})
            continue
        # Style của Word thắng mẫu chữ: đã khai báo rồi thì không đoán lại.
        level = style_level if style_level is not None else para_levels[para_index]
        para_index += 1
        blocks.append({"text": text, "heading_level": level, "page": None})
    return blocks


def _lines_tu_text(text: str) -> list[str]:
    out = []
    for line in (text or "").splitlines():
        t = line.replace("\x00", "").strip()
        if t:
            out.append(t)
    return out


def _trich_mot_bang(plumber_page, bang) -> tuple[list[list], list[list]]:
    """(hàng chế độ mặc định, hàng chế độ text) của MỘT bảng đã dò được.

    Chế độ mặc định lấy THẲNG `bang.extract()` — lưới mà `find_tables()` đã dò
    trên TOÀN trang. KHÔNG dò lại chế độ mặc định bên trong `within_bbox`:
    LỖI THẬT tìm ra 2026-09-04 khi rà toàn nhánh — cắt trang theo đúng bbox của
    bảng làm MẤT các đường kẻ nằm ĐÚNG TRÊN biên cắt, nên `find_tables()` chạy
    lại trong vùng cắt dò ra một lưới NGHÈO HƠN hẳn. Đo trên
    `luat-thuexuatnhapkhau.pdf` trang 13: `bang.extract()` cho 15 hàng x 4 cột
    (STT | Nhóm hàng | Mô tả | Khung thuế suất) còn dò lại trong bbox chỉ cho
    13 hàng x 2 cột — cột "Khung thuế suất" biến mất khỏi lưới, và vì nó nằm
    TRONG bbox nên cũng bị `parse_pdf` cắt khỏi dải văn xuôi: mất hẳn, không
    còn đường nào phát ra. Nhánh dự phòng `bang.extract()` cũ chỉ chạy khi dò
    lại trả về RỖNG, không cứu được ca "dò lại ra lưới nghèo hơn nhưng không
    rỗng" này (trang 16: 25x4 → 23x2).

    Chế độ `text` vẫn phải dò trong `within_bbox` (không có đường nào khác để
    chạy `horizontal_strategy="text"` giới hạn trong một bảng). Dò ra != 1
    bảng thì bỏ chế độ `text` cho vùng này — an toàn hơn đoán sai tương ứng
    bảng nào với bảng nào (spec §3.3 chỉ nói gộp TRONG một bảng, không nói gộp
    bảng CHÉO NHAU)."""
    vung = plumber_page.within_bbox(bang.bbox, relative=False)
    theo_text = vung.find_tables(table_settings={"horizontal_strategy": "text"})
    hang_text = theo_text[0].extract() if len(theo_text) == 1 else []
    return bang.extract(), hang_text


def _khoi_bang(plumber_page, bang, pageno: int
               ) -> tuple[list[dict], list[tuple[str, str]]]:
    """Blocks + warnings của MỘT bảng đã dò được trên trang `pageno`."""
    hang_mac_dinh, hang_text = _trich_mot_bang(plumber_page, bang)
    warnings: list[tuple[str, str]] = []
    # Cổng BẤT ĐỒNG SỐ CỘT — lệch số cột thì BỎ HẲN chế độ `text`:
    # `merge_table_rows(..., [])` rơi đúng vào đường dự phòng sẵn có, chỉ dùng
    # chế độ mặc định.
    bat_dong = bat_dong_so_cot(hang_mac_dinh, hang_text)
    if bat_dong:
        n1, n2 = bat_dong
        hang_text = []
        warnings.append((f"trang {pageno}, bảng",
                         f"hai chế độ trích xuất bất đồng số cột ({n1} vs {n2})"
                         " — dùng riêng chế độ mặc định"))
    gop = merge_table_rows(hang_mac_dinh, hang_text)
    tat_ca_hang = [r for r, _ in gop]
    header_rows, body_rows, _ = split_header_body(tat_ca_hang)
    columns = column_names(header_rows)
    # Bỏ hàng KHÔNG MANG GIÁ TRỊ NÀO — xem `hang_khong_gia_tri`. Im lặng đúng:
    # ô đệm rỗng của biểu mẫu không phải nội dung bị mất.
    blocks = [{"text": row_to_text(row, columns), "heading_level": None,
              "page": pageno, "atomic": True} for row in body_rows
              if not hang_khong_gia_tri(row)]
    gap = checksum_gap(body_rows)
    if gap:
        warnings.append((f"trang {pageno}, bảng", gap))
    return blocks, warnings


def _doc_trang_bang_anh(path: str, pageno: int
                        ) -> tuple[list[str], float | None,
                                   tuple[str, str] | None, list[list[str]]]:
    """Đọc MỘT trang không có lớp text bằng ảnh.

    Trả `(dòng, mean_conf, cảnh_báo, grid)`. Hỏng thì to tiếng NHƯNG không làm
    vỡ cả lượt nạp: cảnh báo mang tên trang đi tiếp qua `IngestReport`, và nếu
    cuối cùng cả tệp không sinh được block nào thì `_ingest_known` đã sẵn ném
    `IngestError` — tệp bị TỪ CHỐI CÓ TÊN, đúng hạ tầng Kế hoạch 1 dựng
    (spec 2026-09-04-tang-ocr §12). Không phát minh cơ chế mới.

    `grid` (bậc 2) chỉ khác rỗng khi lưới dựng được VÀ không hỏng. Bậc 2 hỏng
    (`Region.grid_error` có giá trị) không làm mất `dòng` — nội dung vẫn về
    đủ, chỉ mất cấu trúc cột — nhưng vẫn phải kèm cảnh báo CÓ TÊN (spec §9).
    """
    try:
        kq = read_page(path, pageno)
    except TesseractMissing as e:
        return [], None, (f"trang {pageno}",
                          f"trang không có lớp text và không đọc được bằng ảnh: {e}"), []
    except Exception as e:                      # noqa: BLE001
        # Một trang hỏng (PDF vỡ, ảnh không rasterise được) không được kéo
        # theo cả tài liệu — nhưng phải GỌI TÊN, không nuốt.
        return [], None, (f"trang {pageno}",
                          f"đọc trang bằng ảnh thất bại: {type(e).__name__}: {e}"), []
    lines = _lines_tu_text(kq.text)
    if not lines:
        return [], None, (f"trang {pageno}",
                          "đọc bằng ảnh ra text RỖNG — trang có thể là ảnh trắng "
                          "hoặc bản scan hỏng; KHÔNG nạp gì cho trang này"), []
    grid_error = next((r.grid_error for r in kq.regions if r.grid_error), None)
    if grid_error:
        # Bậc 2 hỏng: nội dung vẫn về đủ (dòng phẳng), chỉ mất cấu trúc cột.
        # Nói TO chứ không nuốt (spec §9) — khác các nhánh hỏng khác ở trên
        # (trả `[]` cho dòng): ở đây `lines` vẫn KHÔNG rỗng, có chủ ý.
        return lines, kq.mean_conf, (
            f"trang {pageno}",
            f"đọc được chữ nhưng KHÔNG dựng được cấu trúc bảng: {grid_error}"), []
    grid = next((r.grid for r in kq.regions if len(r.grid) > 1), [])
    return lines, kq.mean_conf, None, grid


def _khoi_tu_luoi_anh(grid: list[list[str]], pageno: int, furniture: set,
                      conf: float | None) -> list[dict]:
    """Blocks của MỘT trang đọc-từ-ảnh có lưới, GIỮ NGUYÊN thứ tự hàng.

    Bậc 1 nhả ĐÚNG MỘT vùng phủ CẢ TRANG, nên lưới bậc 2 phủ cả letterhead,
    tiêu đề và chân trang chứ không riêng thân bảng. Đưa TRỌN lưới đó vào
    `split_header_body`/`column_names` là biến letterhead thành TÊN CỘT — đo
    được trên SCID tr12 trước khi sửa: tên cột dài 136 ký tự, lặp trong MỌI
    block (phát hiện C1 của review toàn nhánh).

    Nên chia đôi đường đi:
      - hàng nằm trong một DẢI trông như bảng (`table.table_row_runs`) đi
        đường B4: `split_header_body`/`column_names`/`row_to_text` áp trên
        RIÊNG dải đó, mỗi hàng thân một block atomic;
      - hàng ngoài dải nối ô bằng dấu cách thành MỘT dòng rồi đi ĐÚNG đường
        dòng-phẳng cũ — qua `heading_level()` và qua bộ lọc furniture. Việc
        này đóng luôn I4: bản trước, trang có lưới `continue` sớm nên mất CẢ
        HAI (không hàng nào của trang được chấm tiêu đề, không hàng nào bị
        lọc rác đầu/chân trang).
    """
    blocks: list[dict] = []
    runs = dict(table.table_row_runs(grid))
    i = 0
    while i < len(grid):
        if i in runs:
            end = runs[i]
            header_rows, body_rows, _ = split_header_body(grid[i:end])
            # `split_header_body` coi MỌI hàng trước hàng-có-số-thuần đầu
            # tiên là header, nên một token số lạc trong letterhead biến
            # letterhead thành TÊN CỘT. Tìm lại header theo NỘI DUNG trên
            # lưới ĐẦY ĐỦ (không phải lát cắt của dải) và ưu tiên nó khi có.
            # Đo 2026-09-08: tên cột có nghĩa 0,301 -> 0,611 trên 95 trang /
            # 6 tài liệu scan. Xem bảng cạnh `table.MAX_RUN_GAP_ROWS`.
            hr2 = table.find_header_rows(grid, i + len(header_rows))
            columns = column_names(hr2 or header_rows)
            for row in body_rows:
                # Hàng thân KHÔNG mang dữ liệu số là văn xuôi lọt vào dải
                # (tiêu đề mục, câu chú thích, mảnh letterhead, khối chữ ký).
                # Xé nó thành cột làm hỏng chunk, nên đưa về ĐÚNG đường
                # dòng-phẳng như hàng ngoài dải — qua `heading_level()` và
                # qua bộ lọc furniture. Xem `table.has_numeric_data`.
                if not table.has_numeric_data(row):
                    blocks.extend(_flat_line_block(row, pageno, conf,
                                                   furniture))
                    continue
                blocks.append({"text": row_to_text(row, columns,
                                                   compact=True),
                               "heading_level": None, "page": pageno,
                               "atomic": True, "source_kind": "ocr",
                               "ocr_conf": conf})
            i = end
            continue
        blocks.extend(_flat_line_block(grid[i], pageno, conf, furniture))
        i += 1
    return blocks


def _flat_line_block(row: list[str], pageno: int, conf, furniture) -> list[dict]:
    """Một hàng lưới -> ĐÚNG đường dòng-phẳng cũ: nối ô bằng dấu cách, chấm
    `heading_level()`, lọc furniture. Trả [] khi hàng rỗng hoặc là rác đầu/
    chân trang.

    Tách thành hàm riêng 2026-09-08 vì nay có HAI chỗ gọi: hàng nằm ngoài dải
    bảng, và hàng THÂN không mang dữ liệu số."""
    text = " ".join(c.strip() for c in row if c.strip()).strip()
    if not text or _normalize_digits(text) in furniture:
        return []
    return [{"text": text, "heading_level": heading_level(text),
             "page": pageno, "source_kind": "ocr", "ocr_conf": conf}]


def parse_pdf(path: str) -> tuple[list[dict], list[tuple[str, str]]]:
    """Heuristic headings (no font info): numbered/keyword headings & short
    ALL-CAPS lines. Trang CÓ bảng đi qua `pdfplumber` (spec 2026-09-04-b4);
    trang KHÔNG có bảng giữ NGUYÊN đường `pypdf` — bất biến byte-identical
    bắt buộc, xem test `test_trang_khong_bang_byte_identical`.

    HAI LƯỢT từ 2026-08-19: gom dòng theo trang trước, nhận diện rác
    header/footer trên toàn tài liệu, rồi mới dựng block. Một lượt thì không
    thể biết một dòng có lặp trên phần lớn số trang hay không.
    """
    reader = pypdf.PdfReader(path)
    all_warnings: list[tuple[str, str]] = []
    with pdfplumber.open(path) as pdf:
        pages: list[list[str]] = []
        # Đầu vào RIÊNG cho detect_page_furniture — LUÔN pypdf toàn trang, kể
        # cả trang có bảng. `pages` (build block) dùng dải-text pdfplumber cho
        # trang có bảng, khác cách tách dòng của pypdf; nếu detect_page_furniture
        # nhận thẳng `pages`, một dòng furniture thật trích lệch chút giữa hai
        # đường có thể đổi tần suất/vị-trí-rìa của nó và lệch tập furniture của
        # CẢ TÀI LIỆU — kể cả ảnh hưởng tới trang KHÔNG bảng, dù trang đó không
        # đổi gì trong đường trích của chính nó (review Task 2, tài liệu hỗn hợp
        # dạng bieumau_bctc_hopnhat.pdf mới lộ ra, test đồng nhất không bắt được).
        pages_cho_furniture: list[list[str]] = []
        page_bangs: list[list] = []
        ocr_pages: dict[int, float | None] = {}
        grid_by_page: dict[int, list[list[str]]] = {}
        for pageno, page in enumerate(reader.pages, start=1):
            text_toan_trang = _lines_tu_text(page.extract_text() or "")
            # "Rỗng" nghĩa là KHÔNG CÒN DÒNG NÀO sau bước strip — không phải
            # "ít chữ". Máy scan đời mới thường nhúng sẵn một lớp OCR kém nên
            # trang scan có thể trả về vài ký tự rác thay vì rỗng hẳn; KHÔNG
            # đặt ngưỡng "dưới N ký tự" ở đây vì chưa có tài liệu scan thật để
            # hiệu chỉnh N, và hằng số rút từ không khí là thứ dự án này cấm
            # (spec 2026-09-04-tang-ocr §11).
            if not text_toan_trang:
                text_toan_trang, conf, canh_bao, ocr_grid = \
                    _doc_trang_bang_anh(path, pageno)
                if canh_bao:
                    all_warnings.append(canh_bao)
                if text_toan_trang:
                    ocr_pages[pageno] = conf
                    if ocr_grid and max(len(h) for h in ocr_grid) > 1:
                        grid_by_page[pageno] = ocr_grid
            pages_cho_furniture.append(text_toan_trang)
            plumber_page = pdf.pages[pageno - 1]
            bangs = sorted(plumber_page.find_tables(), key=lambda b: b.bbox[1])
            page_bangs.append(bangs)
            if not bangs:
                pages.append(text_toan_trang)
                continue
            if pageno in ocr_pages:
                # Trang này không có lớp text nên đã đi qua OCR ở trên, NHƯNG
                # cũng bị `find_tables()` báo có bảng. Nhánh dưới đây dựng lại
                # dải văn xuôi từ `plumber_page.within_bbox(...).extract_text()`
                # — nguồn hoàn toàn KHÁC `text_toan_trang` (lớp vector, không
                # phải OCR) — nên chữ đọc được bằng ảnh bị BỎ ĐI ở đây mà không
                # ai nói gì: đúng lớp lỗi "mất mát âm thầm" spec này đi đóng.
                # Dựng lại bảng từ ảnh là việc của bậc 2 (chưa có tài liệu scan
                # thật để hiệu chỉnh) — ở bậc 1 chỉ cần GỌI TÊN việc bỏ đi này.
                all_warnings.append((f"trang {pageno}",
                                     "đọc được chữ bằng ảnh nhưng `find_tables()` "
                                     "cũng thấy bảng trên trang — trang này đi "
                                     "đường VECTOR, nên CẢ phần chữ CẢ lưới bảng "
                                     "đọc từ ảnh đều bị bỏ, không vào corpus"))
                # Nhánh dưới đây dựng `dai_lines` từ
                # `plumber_page.within_bbox(...).extract_text()` — LỚP VECTOR,
                # không phải OCR. Nếu KHÔNG bỏ `pageno` khỏi `ocr_pages` ở
                # đây, vòng phát block phía dưới vẫn thấy `pageno in
                # ocr_pages` và dán `source_kind="ocr"` + `mean_conf` của OCR
                # lên text VECTOR — nhãn tin cậy nói dối (review toàn nhánh
                # B4, cùng lớp lỗi spec này đóng).
                #
                # Phải bỏ `grid_by_page[pageno]` CÙNG CHỖ, không chỉ
                # `ocr_pages` (phát hiện I2 của review toàn nhánh): nếu chỉ bỏ
                # `ocr_pages`, vòng phát block phía dưới vẫn thấy `pageno in
                # grid_by_page`, đi nhánh lưới rồi `continue` — nhảy qua CẢ
                # `dai_lines` LẪN `for bang in bangs: _khoi_bang(...)`, tức
                # bảng VECTOR biến mất khỏi corpus trong im lặng, còn cảnh báo
                # ngay trên lại khẳng định điều ngược lại. Kèm theo đó,
                # `ocr_pages.get(pageno)` trả `None` nên block mang
                # `source_kind="ocr"` + `ocr_conf=None` — nhãn tin cậy nói dối
                # lần thứ hai.
                ocr_pages.pop(pageno, None)
                grid_by_page.pop(pageno, None)
            width, height = plumber_page.width, plumber_page.height
            y_bien = [0.0] + [y for b in bangs for y in (b.bbox[1], b.bbox[3])] + [height]
            dai_lines: list[str] = []
            for k in range(0, len(y_bien), 2):
                y0, y1 = y_bien[k], y_bien[k + 1]
                if y1 > y0:
                    dai = plumber_page.within_bbox((0, y0, width, y1), relative=False)
                    dai_lines.extend(_lines_tu_text(dai.extract_text() or ""))
            pages.append(dai_lines)

        furniture = detect_page_furniture(pages_cho_furniture)
        blocks: list[dict] = []
        for pageno, lines in enumerate(pages, start=1):
            plumber_page = pdf.pages[pageno - 1]
            bangs = page_bangs[pageno - 1]
            if pageno in grid_by_page:
                blocks.extend(_khoi_tu_luoi_anh(grid_by_page[pageno], pageno,
                                                furniture,
                                                ocr_pages.get(pageno)))
                continue
            for text in lines:
                if _normalize_digits(text) in furniture:
                    continue
                blk = {"text": text, "heading_level": heading_level(text),
                       "page": pageno}
                if pageno in ocr_pages:
                    blk["source_kind"] = "ocr"
                    blk["ocr_conf"] = ocr_pages[pageno]
                blocks.append(blk)
            for bang in bangs:
                bang_blocks, bang_warnings = _khoi_bang(plumber_page, bang, pageno)
                blocks.extend(bang_blocks)
                all_warnings.extend(bang_warnings)
    return blocks, all_warnings


def _pptx_table_to_text(tbl) -> str:
    """Bảng trong slide thành text nhiều dòng, mỗi dòng một hàng, cột ngăn
    bởi "|" — CÙNG khuôn với `_bang_thanh_text` của .docx, để tầng chunk và
    tầng tổng hợp chỉ phải hiểu một dạng bảng."""
    rows = []
    for row in tbl.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _pptx_shape_type(shape):
    """`shape.shape_type` đọc PHÒNG THỦ — trả None khi không đọc được.

    python-pptx ném `NotImplementedError("Shape instance of unrecognized shape
    type")` với một `<p:sp>` không placeholder, không `prstGeom`, không
    `custGeom` và không cờ `txBox`. Đó là loại shape mà công cụ NGOÀI
    PowerPoint sinh ra — KỂ CẢ LibreOffice, tức chính cầu chuyển đổi
    `.ppt → .pptx` mà nhánh này vừa dựng. Không ai bắt lỗi đó ở phía trên nên
    nó sẽ sập TRỌN `ingest_path` và mất báo cáo của mọi tệp trước đó."""
    try:
        return shape.shape_type
    except NotImplementedError:
        return None


def _pptx_xml_text(shape) -> str:
    """Mọi nút `a:t` trong XML của shape, nối bằng xuống dòng.

    Lưới an toàn cuối cùng cho shape không có `text_frame` mà vẫn mang chữ."""
    el = getattr(shape, "element", None)
    if el is None:
        return ""
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    try:
        parts = [(t.text or "").strip() for t in el.findall(".//a:t", ns)]
    except Exception:
        return ""
    return "\n".join(p for p in parts if p).strip()


def _pptx_shape_blocks(shape, page: int, title_shape) -> list[dict]:
    """Blocks từ MỘT shape của slide, ĐỆ QUY vào group shape.

    Sự thật đo được (python-pptx 1.0.2): một group shape (PowerPoint "Group
    Objects") xuất hiện ở `slide.shapes` như MỘT shape duy nhất, và chính
    shape đó có `has_table = False` VÀ `has_text_frame = False` — bảng/textbox
    thật nằm bên trong `shape.shapes` (con của group), không lộ ra ở cấp
    ngoài. Vòng lặp cũ chỉ kiểm `has_table`/`has_text_frame` ở cấp ngoài nên
    bỏ qua toàn bộ group: nội dung bên trong biến mất KHỎI CORPUS MÀ KHÔNG CÓ
    CẢNH BÁO NÀO — slide vẫn có tiêu đề nên vẫn sinh block, báo cáo nạp vẫn
    "thành công". Group có thể lồng nhiều tầng (group trong group), nên đây
    phải là đệ quy thật — mở đúng MỘT tầng rồi dừng vẫn để lọt group lồng.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if _pptx_shape_type(shape) is MSO_SHAPE_TYPE.GROUP:
        blocks: list[dict] = []
        for sub_shape in getattr(shape, "shapes", ()):
            blocks.extend(_pptx_shape_blocks(sub_shape, page, title_shape))
        return blocks
    if getattr(shape, "has_table", False):
        text = _pptx_table_to_text(shape.table)
        return [{"text": text, "heading_level": None, "page": page}] if text else []
    if getattr(shape, "has_text_frame", False):
        # python-pptx dựng PROXY MỚI mỗi lần truy cập `slide.shapes.title`
        # (đo được: `s.shapes.title is s.shapes.title` → False), nên so danh
        # tính proxy (`shape is title_shape`) KHÔNG BAO GIỜ khớp — tiêu đề lọt
        # qua nhánh này và vào corpus LẦN THỨ HAI (lần đầu ở `parse_pptx` khi
        # sinh heading cấp 1). So phần tử XML nền — `_element` — vì đó là thứ
        # ổn định giữa hai lần truy cập, proxy chỉ bọc quanh nó.
        if title_shape is not None and shape._element is title_shape._element:
            return []
        text = shape.text_frame.text.strip()
        return [{"text": text, "heading_level": None, "page": page}] if text else []

    # Không rơi vào nhánh nào ở trên. KHÔNG được trả [] ngay: đó đúng là cách
    # group shape từng mất im lặng. Quét mọi nút văn bản `a:t` trong XML của
    # shape — bắt được các shape lạ có chữ mà python-pptx không dựng thành
    # text_frame.
    text = _pptx_xml_text(shape)
    if text:
        return [{"text": text, "heading_level": None, "page": page}]

    # Vẫn không có chữ nào. Nếu ĐỌC ĐƯỢC loại shape (ảnh, đường kẻ, media...)
    # thì im lặng là đúng — chúng vốn không mang chữ. Nhưng `shape_type` bằng
    # None nghĩa là python-pptx KHÔNG NHẬN RA shape: với GraphicFrame đó gần
    # như luôn là SmartArt (`GraphicFrame.shape_type` trả None khi graphicData
    # không phải table/chart/OLE). Nội dung SmartArt nằm trong một part sơ đồ
    # riêng mà parser này chưa đọc — nên phải để lại DẤU VẾT QUAN SÁT ĐƯỢC,
    # đúng nguyên tắc "không bao giờ mất im lặng" của spec 2026-08-29 mục 4.
    if _pptx_shape_type(shape) is None:
        name = (getattr(shape, "name", "") or "").strip() or "không tên"
        return [{"text": f"[NỘI DUNG CHƯA ĐỌC ĐƯỢC] shape \"{name}\" trên slide "
                         f"{page} thuộc loại python-pptx không nhận ra (thường "
                         f"là SmartArt). Nội dung của nó KHÔNG có trong corpus.",
                 "heading_level": None, "page": page}]
    return []


def parse_pptx(path: str) -> list[dict]:
    """Blocks theo thứ tự slide; `page` mang SỐ SLIDE (từ 1).

    Tiêu đề slide thành heading cấp 1 — đó là phân cấp duy nhất một bộ slide
    có. Ghi chú thuyết trình được giữ vì trong tài liệu nội bộ chúng thường
    chứa điều kiện và ngoại lệ mà slide chỉ nói tóm tắt.
    """
    from pptx import Presentation

    prs = Presentation(path)
    blocks: list[dict] = []
    for idx, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = (title_shape.text or "").strip() if title_shape is not None else None
        if title:
            blocks.append({"text": title, "heading_level": 1, "page": idx})

        for shape in slide.shapes:
            blocks.extend(_pptx_shape_blocks(shape, idx, title_shape))

        if slide.has_notes_slide:
            note = (slide.notes_slide.notes_text_frame.text or "").strip()
            if note:
                blocks.append({"text": f"Ghi chú: {note}",
                               "heading_level": None, "page": idx})
    return blocks


def _spread_merged(ws) -> list[list]:
    """Lưới giá trị của sheet, với ô gộp được TRẢI ra toàn vùng.

    openpyxl chỉ đặt giá trị ở ô trên-trái của vùng gộp, các ô còn lại là
    None. Không trải thì nhãn cha ("Quý II" trải B:C) chỉ dính vào cột B, và
    cột C mất nhãn — đúng bốn con số 1250/1100/2400/2050 không phân biệt được.
    """
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    for rng in ws.merged_cells.ranges:
        r0, c0, r1, c1 = rng.min_row, rng.min_col, rng.max_row, rng.max_col
        if r0 - 1 >= len(grid) or c0 - 1 >= len(grid[r0 - 1]):
            continue
        value = grid[r0 - 1][c0 - 1]
        if value is None:
            continue
        for r in range(r0 - 1, min(r1, len(grid))):
            for c in range(c0 - 1, min(c1, len(grid[r]))):
                grid[r][c] = value
    return grid


def parse_xlsx(path: str) -> tuple[list[dict], list[tuple[str, str]]]:
    """Trả (sheets, warnings). `warnings` là các cặp (tên sheet, lý do).

    KHÔNG dùng `read_only=True`: chế độ đó không nạp `ws.merged_cells`, mà
    nhãn cột của báo cáo tài chính Việt Nam gần như luôn nằm trong ô gộp.
    Đã thử trên sổ kế toán thật 12,5 MB / 84 sheet: mở được, chấp nhận được.
    """
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    sheets: list[dict] = []
    warnings: list[tuple[str, str]] = []
    for ws in wb.worksheets:
        grid = _spread_merged(ws)          # trải giá trị ô gộp ra toàn vùng
        rows = [r for r in grid if any(c is not None for c in r)]
        if not rows:
            continue
        guess = find_header(rows)
        if guess is None:
            warnings.append((ws.title, "không dò được hàng tiêu đề — "
                                       "nhãn cột để trống thay vì đoán bừa"))
            columns = ["" for _ in rows[0]]
            body = rows
        else:
            parent = rows[guess.row_index - 1] if guess.row_index > 0 else None
            columns = (compose_two_tier(parent, guess.labels)
                       if parent is not None else guess.labels)
            body = rows[guess.row_index + 1:]
        sheets.append({"sheet": ws.title, "columns": columns, "rows": body})
    wb.close()
    return sheets, warnings
