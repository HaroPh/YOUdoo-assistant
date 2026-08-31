# backend/tests/rag/test_parse_pptx.py
"""Parser slide — spec 2026-08-29 mục 5.1.

Fixture tự dựng, đáp án chắc 100% (spec mục 6.1 tầng 1). Chưa có tệp .pptx
thật trong kho test; khi có thì thêm một test `live` đọc nó.
"""
import pytest

pptx = pytest.importorskip("pptx")

from src.rag.parse import parse_pptx


def _make_deck(path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Quy trình bán hàng"
    s1.placeholders[1].text = "Bước 1: tiếp nhận yêu cầu\nBước 2: báo giá"
    s1.notes_slide.notes_text_frame.text = "Nhấn mạnh thời hạn báo giá 24 giờ."

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Định mức chiết khấu"
    tbl = s2.shapes.add_table(3, 2, Inches(1), Inches(2),
                              Inches(6), Inches(2)).table
    tbl.cell(0, 0).text = "Sản lượng"; tbl.cell(0, 1).text = "Chiết khấu"
    tbl.cell(1, 0).text = "dưới 100";  tbl.cell(1, 1).text = "5%"
    tbl.cell(2, 0).text = "từ 100";    tbl.cell(2, 1).text = "10%"
    prs.save(path)


def test_tieu_de_slide_thanh_heading(tmp_path):
    p = str(tmp_path / "deck.pptx"); _make_deck(p)
    blocks = parse_pptx(p)
    heads = [b["text"] for b in blocks if b["heading_level"]]
    assert "Quy trình bán hàng" in heads
    assert "Định mức chiết khấu" in heads


def test_so_slide_duoc_ghi_vao_page(tmp_path):
    p = str(tmp_path / "deck.pptx"); _make_deck(p)
    blocks = parse_pptx(p)
    assert {b["page"] for b in blocks} == {1, 2}


def test_bang_giu_duoc_rang_buoc_hang(tmp_path):
    """Lỗi 4 của spec là số bị tách khỏi nhãn của nó. Bảng trong slide phải
    ra một dòng một hàng, cột ngăn bằng dấu sổ đứng — như `parse_docx`."""
    p = str(tmp_path / "deck.pptx"); _make_deck(p)
    text = "\n".join(b["text"] for b in parse_pptx(p))
    assert "dưới 100 | 5%" in text
    assert "từ 100 | 10%" in text
    assert "Sản lượng | Chiết khấu" in text


def test_ghi_chu_thuyet_trinh_duoc_giu(tmp_path):
    p = str(tmp_path / "deck.pptx"); _make_deck(p)
    text = "\n".join(b["text"] for b in parse_pptx(p))
    assert "24 giờ" in text


def test_slide_rong_khong_sinh_block_rac(tmp_path):
    from pptx import Presentation
    p = str(tmp_path / "trong.pptx")
    prs = Presentation(); prs.slides.add_slide(prs.slide_layouts[6]); prs.save(p)
    assert parse_pptx(p) == []


def test_bang_trong_group_shape_khong_bi_bo_sot(tmp_path):
    """Vòng sửa 1 (reviewer). Group shape của PowerPoint có `has_table=False`
    và `has_text_frame=False` ở chính nó — bảng/textbox thật nằm trong
    `shape.shapes`. Không mở group ra thì cả hai biến mất im lặng, dù slide
    vẫn có tiêu đề nên báo cáo nạp vẫn "thành công"."""
    from pptx import Presentation
    from pptx.util import Inches

    p = str(tmp_path / "group.pptx")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    table_shape = slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    tbl = table_shape.table
    tbl.cell(0, 0).text = "Sản lượng"; tbl.cell(0, 1).text = "Chiết khấu"
    tbl.cell(1, 0).text = "dưới 100";  tbl.cell(1, 1).text = "5%"
    tbl.cell(2, 0).text = "từ 100";    tbl.cell(2, 1).text = "10%"

    textbox = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1))
    textbox.text_frame.text = "Ghi chú nằm trong group"

    slide.shapes.add_group_shape([table_shape, textbox])
    prs.save(p)

    text = "\n".join(b["text"] for b in parse_pptx(p))
    assert "dưới 100 | 5%" in text
    assert "Ghi chú nằm trong group" in text


def test_group_long_nhieu_tang_van_lay_duoc(tmp_path):
    """Bắt lỗi 'chỉ mở một tầng': group chứa group, textbox mang chuỗi nhận
    dạng riêng nằm ở tầng trong cùng. Đệ quy đúng thì xanh; mở một tầng rồi
    dừng thì đỏ."""
    from pptx import Presentation
    from pptx.util import Inches

    p = str(tmp_path / "nested_group.pptx")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    marker = "CHUOI_NHAN_DANG_RIENG_LONG_SAU_HAI_TANG_GROUP"
    inner_textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    inner_textbox.text_frame.text = marker
    inner_group = slide.shapes.add_group_shape([inner_textbox])

    other_textbox = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
    other_textbox.text_frame.text = "shape khác ở tầng ngoài"
    slide.shapes.add_group_shape([inner_group, other_textbox])
    prs.save(p)

    text = "\n".join(b["text"] for b in parse_pptx(p))
    assert marker in text
