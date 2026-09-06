# backend/tests/rag/test_parse_xlsx_merged.py
"""Ô gộp và header hai tầng — spec 2026-08-29 mục 5.2.

Fixture tự dựng, đáp án chắc 100% (spec mục 6.1 tầng 1).
"""
import openpyxl
import pytest

from src.rag.parse import parse_xlsx


def _make_bctc(path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "BCTC"
    ws["A1"] = "BÁO CÁO TÀI CHÍNH HỢP NHẤT QUÝ II/2026"
    ws.merge_cells("A1:E1")
    ws["A3"] = "Chỉ tiêu"; ws["B3"] = "Quý II"; ws["D3"] = "Lũy kế"
    ws.merge_cells("B3:C3"); ws.merge_cells("D3:E3"); ws.merge_cells("A3:A4")
    ws["B4"] = "2026"; ws["C4"] = "2025"; ws["D4"] = "2026"; ws["E4"] = "2025"
    for i, row in enumerate([("Doanh thu thuần", 1250, 1100, 2400, 2050),
                             ("Giá vốn hàng bán", 900, 820, 1750, 1560),
                             ("Lợi nhuận gộp", 350, 280, 650, 490)], start=5):
        for j, v in enumerate(row):
            ws.cell(row=i, column=j + 1, value=v)
    wb.save(path)


def test_nhan_cot_khong_con_la_tieu_de_tai_lieu(tmp_path):
    p = str(tmp_path / "bctc.xlsx"); _make_bctc(p)
    sheets, _ = parse_xlsx(p)
    cols = sheets[0]["columns"]
    assert "BÁO CÁO TÀI CHÍNH" not in " ".join(cols)
    assert cols[0] == "Chỉ tiêu"


def test_header_hai_tang_ghep_cha_con(tmp_path):
    """Bốn con số 1250/1100/2400/2050 chỉ phân biệt được nhờ nhãn ghép."""
    p = str(tmp_path / "bctc.xlsx"); _make_bctc(p)
    sheets, _ = parse_xlsx(p)
    cols = sheets[0]["columns"]
    assert "Quý II · 2026" in cols
    assert "Quý II · 2025" in cols
    assert "Lũy kế · 2026" in cols
    assert "Lũy kế · 2025" in cols


def test_o_gop_duoc_trai_gia_tri_khong_con_None(tmp_path):
    p = str(tmp_path / "bctc.xlsx"); _make_bctc(p)
    sheets, _ = parse_xlsx(p)
    assert all(c for c in sheets[0]["columns"][:5])


def test_hang_du_lieu_khong_lan_hang_tieu_de(tmp_path):
    p = str(tmp_path / "bctc.xlsx"); _make_bctc(p)
    sheets, _ = parse_xlsx(p)
    first = sheets[0]["rows"][0]
    assert first[0] == "Doanh thu thuần"


def test_sheet_khong_do_duoc_tieu_de_thi_sinh_canh_bao(tmp_path):
    p = str(tmp_path / "soden.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SoLieu"
    for i, row in enumerate([[1, 2, 3], [4, 5, 6], [7, 8, 9]], start=1):
        for j, v in enumerate(row):
            ws.cell(row=i, column=j + 1, value=v)
    wb.save(p)
    _, canh_bao = parse_xlsx(p)
    assert any(sheet == "SoLieu" for sheet, _ in canh_bao)
