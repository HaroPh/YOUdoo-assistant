# backend/tests/rag/test_ingest_run.py
"""Mức LƯỢT CHẠY — `ingest_path()` và `main()` (spec 2026-08-29 mục 4).

Vì sao có tệp này (review toàn nhánh 2026-08-31):

* `main()` chưa có một test nào ở BẤT KỲ tầng nào. Nhánh `os.walk`, `render()`
  in ra, và `sys.exit(1)` chưa từng chạy — gỡ hẳn `sys.exit(1)` thì 2242 test
  vẫn xanh. Đó đúng là lời hứa nổi bật nhất của spec mà không ai đo.
* `ingest_path` không có lá chắn quanh thân vòng lặp, nên MỘT tệp hỏng làm sập
  trọn lượt nạp và xoá sạch báo cáo của mọi tệp trước đó.
* `ingest_path` trỏ vào đường dẫn không tồn tại trả `IngestReport(0,0,0,[])`
  với `ok=True`: gõ sai đường dẫn cho ra một BÁO CÁO THÀNH CÔNG.

Không cần Postgres: mọi lối ra DB đều thay bằng `_FakeConn`.
"""
import contextlib

import pytest

from src.rag import ingest as _ing


class _FakeConn:
    """Chưa từng nạp tệp nào; nuốt mọi lệnh ghi; đủ cho `with conn.transaction()`."""

    def __init__(self):
        self.closed = False

    def execute(self, *a, **k):
        return self

    def fetchone(self):
        return None

    def transaction(self):
        return contextlib.nullcontext()

    def close(self):
        self.closed = True


class _FakeEmbedder:
    model_name = "gia-lap"
    dim = 2


@pytest.fixture
def _khong_cham_db_va_embed(monkeypatch):
    """Chặn cả ba lối ra ngoài tiến trình: kết nối DB, dựng schema, embed."""
    monkeypatch.setattr(_ing, "get_embedder", lambda: _FakeEmbedder())
    monkeypatch.setattr(_ing, "embed_texts", lambda texts: [[0.0, 0.0] for _ in texts])
    monkeypatch.setattr(_ing._db, "connect", lambda *a, **k: _FakeConn())
    monkeypatch.setattr(_ing._db, "ensure_schema", lambda *a, **k: None)


def _docx_that(path, tieu_de="Điều 1", than="Nội dung thử nghiệm."):
    from docx import Document
    d = Document()
    d.add_heading(tieu_de, level=1)
    d.add_paragraph(than)
    d.save(str(path))
    return path


# ----------------------------------------------------------------- C4

def test_mot_tep_hong_KHONG_lam_mat_bao_cao_cua_tep_sau(_khong_cham_db_va_embed,
                                                        tmp_path):
    """C4: mỗi tệp phải ra ở đúng một trong ba trạng thái, kể cả khi tệp ngay
    trước nó ném ngoại lệ.

    Tệp hỏng ở đây là một `.docx` KHÔNG phải gói OOXML — `parse_docx` ném
    `PackageNotFoundError` thật, không giả lập gì. Cùng cửa đó còn thoát ra
    `IngestError` (PDF không lớp text), `InvalidFileException`/`BadZipFile`
    (.xlsx), `NotImplementedError` (.pptx của công cụ ngoài PowerPoint).

    `01_` / `02_` để `os.walk` trả tệp hỏng TRƯỚC: nếu tệp tốt chạy trước thì
    lá chắn không được đo (báo cáo vẫn có `ingested`)."""
    (tmp_path / "01_hong.docx").write_bytes(b"khong phai OOXML")
    _docx_that(tmp_path / "02_tot.docx")

    rep = _ing.ingest_path(str(tmp_path), conn=_FakeConn())

    assert rep.ingested == 1, "tệp lành phía sau tệp hỏng phải vẫn được nạp"
    assert len(rep.rejected) == 1
    assert rep.rejected[0].path.endswith("01_hong.docx")
    # Lý do phải mang LOẠI lỗi và thông điệp, đủ để người đọc biết chuyện gì.
    assert "PackageNotFoundError" in rep.rejected[0].reason
    assert rep.ok is False


def test_PDF_khong_lop_text_ra_rejected_CO_TEN_chu_khong_sap_luot(
        _khong_cham_db_va_embed, tmp_path, monkeypatch):
    """C4, ca được spec mục 4 nêu ĐÍCH DANH: "PDF không lớp text, parse ra
    rỗng" phải là `rejected` có tên. Trước bản sửa nó `raise IngestError` và
    không ai bắt."""
    monkeypatch.setattr(_ing, "_chunks_for", lambda *a, **k: [])
    (tmp_path / "scan.pdf").write_bytes(b"%PDF-1.4 fake")

    rep = _ing.ingest_path(str(tmp_path), conn=_FakeConn())

    assert len(rep.rejected) == 1
    assert rep.rejected[0].path.endswith("scan.pdf")
    assert "IngestError" in rep.rejected[0].reason
    assert "scan" in rep.rejected[0].reason


# ----------------------------------------------------------------- C6

def test_duong_dan_dich_KHONG_TON_TAI_thi_hong_TO_TIENG(_khong_cham_db_va_embed):
    """C6: trước bản sửa, gõ sai đường dẫn trả `IngestReport(0,0,0,[])`,
    `ok=True`, `main()` thoát 0 — một BÁO CÁO THÀNH CÔNG cho một lượt chạy
    không hề chạm được gì. Chính hình dạng đó đã được `test_ingest_kho_that`
    gọi tên là IM_LANG khi nó xảy ra với một TỆP."""
    with pytest.raises(_ing.IngestTargetMissing) as e:
        _ing.ingest_path("d:/khong/he/ton/tai/duong/nay", conn=_FakeConn())
    assert "khong/he/ton/tai" in str(e.value)


def test_thu_muc_CO_THAT_nhung_rong_thi_KHONG_phai_loi(_khong_cham_db_va_embed,
                                                       tmp_path):
    """Mặt kia của C6: thư mục rỗng là HỢP LỆ, không được ném. Nhưng báo cáo
    phải nói rõ là không thấy tài liệu nào — dòng số đếm toàn 0 đọc y hệt một
    lượt nạp đã xong."""
    rep = _ing.ingest_path(str(tmp_path), conn=_FakeConn())
    assert rep.ok is True
    assert (rep.ingested, rep.unchanged, rep.chunks) == (0, 0, 0)
    assert "KHÔNG THẤY TÀI LIỆU NÀO" in rep.render()


def test_thu_muc_chi_co_tep_khong_phai_tai_lieu_van_noi_ro(_khong_cham_db_va_embed,
                                                           tmp_path):
    """`.txt` không phải tài liệu nên không bị từ chối — nhưng lượt chạy vẫn
    không nạp được gì, và người đọc phải thấy điều đó."""
    (tmp_path / "ghi_chu.txt").write_text("khong phai tai lieu", encoding="utf-8")
    rep = _ing.ingest_path(str(tmp_path), conn=_FakeConn())
    assert rep.ok is True
    assert "KHÔNG THẤY TÀI LIỆU NÀO" in rep.render()


# ----------------------------------------------------------------- C5

def test_main_chay_that_tren_thu_muc_hon_hop(_khong_cham_db_va_embed,
                                             monkeypatch, tmp_path, capsys):
    """C5: `main()` chưa từng được gọi bởi bất kỳ test nào. Bài này chạy
    `main()` THẬT trên một thư mục hỗn hợp và khoá cả ba lời hứa của spec
    mục 4 cùng lúc:

        (1) thoát khác 0 khi có tài liệu bị từ chối,
        (2) GỌI TÊN tệp bị từ chối trên stdout,
        (3) tệp KHÔNG phải tài liệu (`.txt`) không bị tính là từ chối —
            nếu không, một `.gitkeep` cũng làm tiến trình thoát khác 0.

    Giả lập KHÔNG có LibreOffice để `.doc` bị từ chối trên mọi máy chạy."""
    from src.rag import convert as _conv
    monkeypatch.setattr(_conv, "soffice_path", lambda: None)

    _docx_that(tmp_path / "quy_che_tot.docx")
    (tmp_path / "quy_che_cu.doc").write_bytes(b"\xd0\xcf\x11\xe0 fake OLE")
    (tmp_path / "ghi_chu.txt").write_text("khong phai tai lieu", encoding="utf-8")

    monkeypatch.setattr(_ing.sys, "argv", ["ingest", str(tmp_path)])

    with pytest.raises(SystemExit) as e:
        _ing.main()
    assert e.value.code == 1                       # (1)

    out = capsys.readouterr().out
    assert "quy_che_cu.doc" in out                 # (2)
    dong_tu_choi = [d for d in out.splitlines() if "TỪ CHỐI" in d]
    assert len(dong_tu_choi) == 1, dong_tu_choi
    assert "ghi_chu.txt" not in out                # (3)
    assert "đã nạp 1" in out


def test_main_thoat_khac_0_khi_duong_dan_khong_ton_tai(_khong_cham_db_va_embed,
                                                       monkeypatch, capsys):
    """C6 ở mức tiến trình: gõ sai đường dẫn phải thoát khác 0 VÀ in một dòng
    đọc được, không phải traceback."""
    monkeypatch.setattr(_ing.sys, "argv", ["ingest", "d:/khong/he/ton/tai/xyz"])
    with pytest.raises(SystemExit) as e:
        _ing.main()
    assert e.value.code == 1
    out = capsys.readouterr().out
    assert "không tồn tại" in out
    assert "xyz" in out


def test_main_thoat_0_khi_moi_tep_deu_nap_duoc(_khong_cham_db_va_embed,
                                               monkeypatch, tmp_path, capsys):
    """Chân đối chứng: `sys.exit(1)` phải PHÂN BIỆT được hai lượt chạy. Không
    có bài này thì một `sys.exit(1)` vô điều kiện vẫn qua được test trên."""
    _docx_that(tmp_path / "a.docx")
    _docx_that(tmp_path / "b.docx", tieu_de="Điều 2", than="Nội dung khác.")
    monkeypatch.setattr(_ing.sys, "argv", ["ingest", str(tmp_path)])

    _ing.main()                                     # KHÔNG được SystemExit
    out = capsys.readouterr().out
    assert "đã nạp 2" in out
    assert "TỪ CHỐI" not in out


def test_embedder_chet_thi_KHONG_bien_thanh_117_tep_bi_tu_choi(
        _khong_cham_db_va_embed, monkeypatch, tmp_path):
    """Ranh giới của lá chắn C4. Lá chắn biến ngoại lệ của MỘT TỆP thành
    `rejected`; nó KHÔNG được nuốt hỏng hạ tầng.

    Embedder chết mà bị tính thành `rejected` sẽ cho ra báo cáo "mọi tài liệu
    đều bị từ chối" trong khi sự thật là "Ollama không chạy" — che đúng
    nguyên nhân và đổ lỗi cho tài liệu của người dùng. Nó cũng làm hỏng đảm
    bảo "không ghi dở": lượt nạp phải dừng ở tệp đầu tiên, không cày hết kho."""
    from src.rag.embed import EmbeddingError

    def _no(texts):
        raise EmbeddingError("ollama down")
    monkeypatch.setattr(_ing, "embed_texts", _no)

    _docx_that(tmp_path / "a.docx")
    _docx_that(tmp_path / "b.docx", tieu_de="Điều 2", than="Nội dung khác.")

    with pytest.raises(EmbeddingError):
        _ing.ingest_path(str(tmp_path), conn=_FakeConn())
