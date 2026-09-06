# backend/tests/rag/test_parse_docx_phan_cap.py
"""`parse_docx` hai lượt — spec 2026-09-04 mục 7.

Fixture tự dựng bằng python-docx, đáp án chắc 100% (tầng 1 của spec
2026-08-29 mục 6.1).
"""
from docx import Document

from src.rag.parse import DOCX_LEVEL, parse_docx


def _levels(path):
    return [b["heading_level"] for b in parse_docx(str(path))]


def test_tai_lieu_KHONG_style_van_suy_duoc_phan_cap(tmp_path):
    """Ca thật của 10 biểu mẫu BCTC: 0 đoạn nào dùng style Heading."""
    p = tmp_path / "khong_style.docx"
    d = Document()
    for line in ("BÁO CÁO THỬ NGHIỆM", "I. Phần thứ nhất",
                 "Nội dung của phần thứ nhất.", "II. Phần thứ hai",
                 "Nội dung của phần thứ hai."):
        d.add_paragraph(line)
    d.save(str(p))

    levels = _levels(p)
    assert levels[0] == DOCX_LEVEL["upper"]
    assert levels[1] == DOCX_LEVEL["roman"]
    assert levels[2] is None
    assert levels[3] == DOCX_LEVEL["roman"]
    assert levels[4] is None


def test_style_Heading_duoc_ANH_XA_sang_cung_thang(tmp_path):
    """`Heading N` → `N * STYLE_SCALE`. Dùng cấp style THÔ là sai: `Heading 2`
    thành cấp 2, cao hơn cả `phan` (5) lẫn `chuong` (10), hất sạch mọi thứ."""
    p = tmp_path / "co_style.docx"
    d = Document()
    d.add_paragraph("Tiêu đề lớn", style="Heading 1")
    d.add_paragraph("Tiêu đề nhỏ", style="Heading 3")
    d.save(str(p))

    levels = _levels(p)
    assert levels == [DOCX_LEVEL["chuong"], DOCX_LEVEL["muc"]]


def test_style_duoc_UU_TIEN_hon_mau_chu(tmp_path):
    """Word đã khai báo rồi thì không đoán lại — không đổi hành vi cũ."""
    p = tmp_path / "tron.docx"
    d = Document()
    d.add_paragraph("I. Dòng này có style", style="Heading 1")
    d.add_paragraph("II. Dòng này không có style")
    d.save(str(p))

    levels = _levels(p)
    assert levels[0] == DOCX_LEVEL["chuong"], "style thắng"
    assert levels[1] == DOCX_LEVEL["roman"], "mẫu chữ lo phần còn lại"


def test_bang_van_la_THAN_khong_bao_gio_la_tieu_de(tmp_path):
    p = tmp_path / "co_bang.docx"
    d = Document()
    d.add_paragraph("I. Mục một")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "1. Cột này trông như đánh số"
    table.cell(0, 1).text = "Giá"
    table.cell(1, 0).text = "2. Hàng nữa"
    table.cell(1, 1).text = "100"
    d.add_paragraph("II. Mục hai")
    d.save(str(p))

    blocks = parse_docx(str(p))
    table_blocks = [b for b in blocks if "|" in b["text"]]
    assert table_blocks, "phải có block bảng"
    assert all(b["heading_level"] is None for b in table_blocks)


def test_thu_tu_block_khong_doi(tmp_path):
    """Hai lượt không được xáo thứ tự — `chunk_text_blocks` dựng breadcrumb
    theo THỨ TỰ block, xáo là gắn nội dung vào mục sai."""
    p = tmp_path / "thu_tu.docx"
    d = Document()
    for line in ("I. Một", "thân một", "II. Hai", "thân hai"):
        d.add_paragraph(line)
    d.save(str(p))

    texts = [b["text"] for b in parse_docx(str(p))]
    assert texts == ["I. Một", "thân một", "II. Hai", "thân hai"]
