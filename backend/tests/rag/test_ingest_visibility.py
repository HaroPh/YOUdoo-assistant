# backend/tests/rag/test_ingest_visibility.py
"""Ingest phải ghi lớp visibility từ DOC_VISIBILITY — lần nạp seed/ sau không
reset về 'all'. Không cần Postgres: conn giả ghi lại INSERT."""
import contextlib
import re

import pytest

from src.rag import ingest as _ing


class _RecConn:
    def __init__(self):
        self.calls: list[tuple[str, tuple]] = []

    def execute(self, sql, params=None):
        self.calls.append((sql, params))
        return self

    def fetchone(self):
        return None            # chưa từng nạp → đi đường INSERT

    def transaction(self):
        return contextlib.nullcontext()


class _FakeEmbedder:
    model_name = "gia-lap"
    dim = 2


@pytest.fixture
def _khong_ra_ngoai(monkeypatch):
    monkeypatch.setattr(_ing, "get_embedder", lambda: _FakeEmbedder())
    monkeypatch.setattr(_ing, "embed_texts", lambda texts: [[0.0, 0.0] for _ in texts])
    monkeypatch.setattr(_ing._db, "ensure_schema", lambda *a, **k: None)


def _docx(path, than="Chiết khấu bậc 5%."):
    from docx import Document
    d = Document()
    d.add_heading("Điều 1", level=1)
    d.add_paragraph(than)
    d.save(str(path))
    return path


def _visibility_da_ghi(conn: _RecConn) -> set[str]:
    """Đọc giá trị visibility từ INSERT rag_chunks theo VỊ TRÍ cột trong SQL —
    không đoán chỉ số cứng."""
    ra = set()
    for sql, params in conn.calls:
        if not sql.lstrip().startswith("INSERT INTO rag_chunks"):
            continue
        cols = re.search(r"INSERT INTO rag_chunks \((.*?)\)", sql, re.S).group(1)
        names = [c.strip() for c in cols.split(",")]
        assert "visibility" in names, sql
        ra.add(params[names.index("visibility")])
    assert ra, "không có INSERT rag_chunks nào — test tự vô hiệu"
    return ra


@pytest.mark.parametrize("ten,lop", [("discount_policy.docx", "commercial"),
                                     ("policy.docx", "all")])
def test_ingest_ghi_dung_lop(_khong_ra_ngoai, tmp_path, ten, lop):
    conn = _RecConn()
    _ing._ingest_file(str(_docx(tmp_path / ten)), conn)
    assert _visibility_da_ghi(conn) == {lop}
