"""Cảnh báo mức sheet đi được tới báo cáo nạp — spec mục 4.

Một sheet không dò được tiêu đề vẫn cho tệp nạp thành công, nhưng phải
được GỌI TÊN. Nuốt nó là dựng lại lỗi mà cả spec đi đóng.
"""
import contextlib

import openpyxl
import pytest

from src.rag import ingest as _ing


class _FakeConn:
    def execute(self, *a, **k):
        return self

    def fetchone(self):
        return None

    def transaction(self):
        return contextlib.nullcontext()


@pytest.fixture
def _no_embed(monkeypatch):
    monkeypatch.setattr(_ing, "embed_texts",
                        lambda texts: [[0.01] * 1024 for _ in texts])


def _make_sheet_khong_tieu_de(path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SoLieu"
    for i, row in enumerate([[1, 2, 3], [4, 5, 6], [7, 8, 9]], start=1):
        for j, v in enumerate(row):
            ws.cell(row=i, column=j + 1, value=v)
    wb.save(path)


def test_sheet_dang_ngo_duoc_goi_ten_trong_bao_cao(tmp_path, _no_embed):
    p = str(tmp_path / "so.xlsx"); _make_sheet_khong_tieu_de(p)
    rep = _ing._ingest_file(p, conn=_FakeConn())
    assert rep.ingested == 1, "tệp vẫn phải nạp được"
    assert rep.ok is True, "cảnh báo KHÔNG làm lượt nạp thất bại"
    assert any(w.where == "SoLieu" for w in rep.warnings)
    assert "so.xlsx" in rep.render()
    assert "SoLieu" in rep.render()


def test_sheet_binh_thuong_khong_sinh_canh_bao(tmp_path, _no_embed):
    """Chân đối chứng: nếu mọi sheet đều sinh cảnh báo thì cảnh báo vô nghĩa."""
    p = str(tmp_path / "sach.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "BangGia"
    ws.append(["Sản phẩm", "Giá", "Tồn"])
    ws.append(["Bàn", 1200000, 30])
    ws.append(["Ghế", 450000, 120])
    wb.save(p)
    rep = _ing._ingest_file(p, conn=_FakeConn())
    assert rep.ingested == 1
    assert rep.warnings == []
